"""
DOCX exporter for FCT proposal pipeline outputs.
Produces three document types: final proposal, proposal summary, and final review.
"""

from __future__ import annotations

import logging
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from src.config.fct_constants import CHAR_LIMITS as CL
from src.generators.models import ConsensusReport, Proposal

logger = logging.getLogger(__name__)


class DocxExporter:
    """Exports pipeline results as professionally formatted DOCX documents."""

    @staticmethod
    def _setup_doc() -> Document:
        """Create a new Document configured for A4, Arial 11pt, 2cm margins."""
        doc = Document()
        section = doc.sections[0]
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.orientation = WD_ORIENT.PORTRAIT
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

        # Set default font
        style = doc.styles["Normal"]
        font = style.font
        font.name = "Arial"
        font.size = Pt(11)
        font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        # Heading styles
        for level in range(1, 4):
            heading_style = doc.styles[f"Heading {level}"]
            heading_style.font.name = "Arial"
            heading_style.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

        return doc

    @staticmethod
    def _add_heading_with_chars(
        doc: Document, text: str, actual: int, limit: int, level: int = 1,
    ) -> None:
        """Add a heading with character count display."""
        doc.add_heading(f"{text} [{actual:,} / {limit:,} chars]", level=level)

    @staticmethod
    def _add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
        """Add a formatted table to the document."""
        table = doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Light Grid Accent 1"

        # Header row
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(10)

        # Data rows
        for row_data in rows:
            row = table.add_row()
            for i, value in enumerate(row_data):
                row.cells[i].text = str(value)
                for paragraph in row.cells[i].paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(10)

        doc.add_paragraph("")  # spacing

    @staticmethod
    def export_final_proposal(proposal: Proposal, path: Path) -> None:
        """Export a submission-ready proposal DOCX.

        Contains all FCT form sections with character counts. No review content.
        """
        doc = DocxExporter._setup_doc()

        # Title page
        title_para = doc.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_para.add_run(proposal.title_en)
        title_run.bold = True
        title_run.font.size = Pt(18)
        title_run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

        doc.add_paragraph("")
        meta_para = doc.add_paragraph()
        meta_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta_para.add_run(
            f"Acronym: {proposal.acronym}  |  Typology: {proposal.typology.value}  |  "
            f"Duration: {proposal.duration_months} months  |  "
            f"Budget: \u20ac{proposal.total_budget:,.0f}"
        ).font.size = Pt(12)

        if proposal.keywords_en:
            kw_para = doc.add_paragraph()
            kw_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            kw_para.add_run(
                f"Keywords: {', '.join(proposal.keywords_en)}"
            ).font.size = Pt(11)

        doc.add_paragraph("")

        # Character count summary table
        doc.add_heading("Character Count Summary", level=2)
        report = proposal.char_count_report()
        rows = []
        for name, counts in report.items():
            status = "OK" if counts["remaining"] >= 0 else "OVER"
            rows.append([
                name.replace("_", " ").title(),
                f"{counts['actual']:,}",
                f"{counts['limit']:,}",
                f"{counts['remaining']:,}",
                status,
            ])
        DocxExporter._add_table(
            doc, ["Section", "Used", "Limit", "Remaining", "Status"], rows,
        )

        doc.add_page_break()

        # Sections with char counts
        sections = [
            ("Abstract (EN)", proposal.abstract_en, CL.abstract_en),
            ("Abstract (PT)", proposal.abstract_pt, CL.abstract_pt),
            ("State of the Art & Objectives",
             proposal.state_of_art_objectives, CL.state_of_art_objectives),
            ("Research Plan & Methods",
             proposal.research_plan_methods, CL.research_plan_methods),
            ("Bibliographic References",
             proposal.bibliographic_references, CL.bibliographic_references),
        ]
        for title, body, limit in sections:
            if not body:
                continue
            DocxExporter._add_heading_with_chars(doc, title, len(body), limit)
            doc.add_paragraph(body)

        # Tasks
        if proposal.tasks:
            doc.add_heading("Tasks", level=1)
            for t in proposal.tasks:
                doc.add_heading(
                    f"T{t.number}: {t.denomination} "
                    f"(months {t.start_month}\u2013"
                    f"{t.start_month + t.duration_months - 1}, "
                    f"{t.person_months} PM)",
                    level=2,
                )
                DocxExporter._add_heading_with_chars(
                    doc, "Description", len(t.description),
                    CL.task_description, level=3,
                )
                doc.add_paragraph(t.description)
                if t.cost_justification:
                    DocxExporter._add_heading_with_chars(
                        doc, "Cost Justification", len(t.cost_justification),
                        CL.cost_justification, level=3,
                    )
                    doc.add_paragraph(t.cost_justification)

        # Deliverables
        if proposal.deliverables:
            doc.add_heading("Deliverables", level=1)
            rows = [
                [d.code, d.title, str(d.due_month), d.description[:200]]
                for d in proposal.deliverables
            ]
            DocxExporter._add_table(
                doc, ["Code", "Title", "Due Month", "Description"], rows,
            )

        # Milestones
        if proposal.milestones:
            doc.add_heading("Milestones", level=1)
            rows = [
                [m.code, m.denomination, str(m.due_month), m.description[:200]]
                for m in proposal.milestones
            ]
            DocxExporter._add_table(
                doc, ["Code", "Denomination", "Due Month", "Description"], rows,
            )

        # Remaining sections
        remaining = [
            ("Management Structure", proposal.management_structure, CL.management_structure),
            ("Institution Description", proposal.institution_description, CL.institution_description),
            ("Career Profile", proposal.career_profile, CL.career_profile),
            ("Contributions \u2014 New Ideas", proposal.contributions_new_ideas, CL.contributions_new_ideas),
            ("Contributions \u2014 Teams", proposal.contributions_teams, CL.contributions_teams),
            ("Contributions \u2014 Society", proposal.contributions_society, CL.contributions_society),
            ("Further Details", proposal.further_details, CL.further_details),
            ("Team CV Synopsis", proposal.team_cv_synopsis, CL.team_cv_synopsis),
            ("Ethics Justification", proposal.ethics_justification, CL.ethics_justification),
        ]
        for title, body, limit in remaining:
            if not body:
                continue
            DocxExporter._add_heading_with_chars(doc, title, len(body), limit)
            doc.add_paragraph(body)

        if proposal.why_timely_pex:
            DocxExporter._add_heading_with_chars(
                doc, "Why Timely (PEX)", len(proposal.why_timely_pex), CL.why_timely_pex,
            )
            doc.add_paragraph(proposal.why_timely_pex)

        doc.save(str(path))
        logger.info(f"  Saved: {path}")

    @staticmethod
    def export_proposal_summary(
        proposal: Proposal,
        consensus: ConsensusReport | None,
        history: list[dict],
        path: Path,
    ) -> None:
        """Export a proposal summary DOCX with key metrics and review highlights."""
        doc = DocxExporter._setup_doc()

        doc.add_heading(proposal.title_en, level=0)

        # Key metrics table
        doc.add_heading("Key Metrics", level=1)
        DocxExporter._add_table(
            doc,
            ["Property", "Value"],
            [
                ["Acronym", proposal.acronym],
                ["Typology", proposal.typology.value],
                ["Duration", f"{proposal.duration_months} months"],
                ["Budget", f"\u20ac{proposal.total_budget:,.0f}"],
                ["Keywords", ", ".join(proposal.keywords_en)],
                ["Scientific Domain", proposal.scientific_domain],
            ],
        )

        # Iteration history table
        if history:
            doc.add_heading("Iteration History", level=1)
            rows = [
                [
                    str(h["iteration"]),
                    f"{h['score']:.1f}",
                    h["decision"],
                    f"{h.get('weighted_scores', {}).get('A', 0):.1f}",
                    f"{h.get('weighted_scores', {}).get('B', 0):.1f}",
                    f"{h.get('weighted_scores', {}).get('C', 0):.1f}",
                ]
                for h in history
            ]
            DocxExporter._add_table(
                doc, ["Iter", "Score", "Decision", "A", "B", "C"], rows,
            )

        # Abstract
        if proposal.abstract_en:
            doc.add_heading("Abstract", level=1)
            doc.add_paragraph(proposal.abstract_en)

        # Abbreviated sections
        for title, body in [
            ("State of the Art & Objectives", proposal.state_of_art_objectives),
            ("Research Plan & Methods", proposal.research_plan_methods),
        ]:
            if body:
                doc.add_heading(title, level=1)
                # Show first 2000 chars for summary
                text = body if len(body) <= 2000 else body[:2000] + "\n\n[... truncated for summary]"
                doc.add_paragraph(text)

        # Panel narrative
        if consensus and consensus.panel_narrative:
            doc.add_heading(
                f"Panel Review (Score: {consensus.consensus_score:.1f}/10 "
                f"\u2014 {consensus.panel_decision})",
                level=1,
            )
            doc.add_paragraph(consensus.panel_narrative)

        doc.save(str(path))
        logger.info(f"  Saved: {path}")

    @staticmethod
    def export_final_review(
        consensus: ConsensusReport | None,
        history: list[dict],
        draft_changes: list[dict],
        path: Path,
    ) -> None:
        """Export a comprehensive review DOCX with scores, history, and draft evolution."""
        if consensus is None:
            return

        doc = DocxExporter._setup_doc()

        doc.add_heading("Final Review Report", level=0)

        # Final scores table
        doc.add_heading("Final Scores", level=1)
        ws = consensus.weighted_scores
        DocxExporter._add_table(
            doc,
            ["Criterion", "Score", "Weight"],
            [
                ["A \u2014 Scientific Merit & Innovation", f"{ws.get('A', 0):.1f}", "40%"],
                ["B \u2014 PI & Team", f"{ws.get('B', 0):.1f}", "30%"],
                ["C \u2014 Feasibility", f"{ws.get('C', 0):.1f}", "30%"],
                ["OVERALL", f"{consensus.consensus_score:.1f}", "100%"],
            ],
        )

        p = doc.add_paragraph()
        run = p.add_run(f"Panel Decision: {consensus.panel_decision.upper()}")
        run.bold = True
        run.font.size = Pt(13)

        # Key findings
        if consensus.key_strengths:
            doc.add_heading("Key Strengths", level=2)
            for s in consensus.key_strengths:
                doc.add_paragraph(s, style="List Bullet")

        if consensus.key_weaknesses:
            doc.add_heading("Key Weaknesses", level=2)
            for w in consensus.key_weaknesses:
                doc.add_paragraph(w, style="List Bullet")

        if consensus.priority_revisions:
            doc.add_heading("Priority Revisions", level=2)
            for i, r in enumerate(consensus.priority_revisions, 1):
                doc.add_paragraph(f"{i}. {r}")

        # Panel narrative
        if consensus.panel_narrative:
            doc.add_heading("Panel Narrative", level=1)
            doc.add_paragraph(consensus.panel_narrative)

        # Review history per iteration
        if history:
            doc.add_heading("Score Progression", level=1)
            rows = [
                [
                    str(h["iteration"]),
                    f"{h['score']:.1f}",
                    h["decision"],
                    f"{h.get('weighted_scores', {}).get('A', 0):.1f}",
                    f"{h.get('weighted_scores', {}).get('B', 0):.1f}",
                    f"{h.get('weighted_scores', {}).get('C', 0):.1f}",
                ]
                for h in history
            ]
            DocxExporter._add_table(
                doc, ["Iter", "Score", "Decision", "A", "B", "C"], rows,
            )

        # Individual reviewer reports
        doc.add_heading("Individual Reviewer Reports", level=1)
        for review in consensus.individual_reviews:
            doc.add_heading(
                f"{review.reviewer_id} ({review.perspective})", level=2,
            )
            doc.add_paragraph(
                f"Model: {review.reviewer_provider}/{review.reviewer_model}  |  "
                f"Score: {review.overall_score:.1f}/10  |  "
                f"Decision: {review.decision}"
            )

            if review.criterion_scores:
                rows = [
                    [
                        cs.sub_criterion or cs.criterion,
                        f"{cs.score:.1f}",
                        cs.justification[:200] if cs.justification else "",
                    ]
                    for cs in review.criterion_scores
                ]
                DocxExporter._add_table(
                    doc, ["Criterion", "Score", "Justification"], rows,
                )

            if review.general_comments:
                doc.add_paragraph(review.general_comments)

            if review.major_revisions:
                doc.add_heading("Major Revisions", level=3)
                for r in review.major_revisions:
                    doc.add_paragraph(r, style="List Bullet")

        # Draft evolution log
        if draft_changes:
            doc.add_heading("Draft Evolution Log", level=1)
            for entry in draft_changes:
                doc.add_heading(
                    f"v{entry['from_version']} \u2192 v{entry['to_version']}", level=2,
                )
                for change in entry.get("changes", []):
                    field = change.get("field", "unknown")
                    review_comment = change.get("review_comment", "")
                    change_desc = change.get("change_description", "")
                    doc.add_paragraph(f"Field: {field}", style="List Bullet")
                    if review_comment:
                        doc.add_paragraph(f"  Review said: {review_comment}")
                    if change_desc:
                        doc.add_paragraph(f"  Changed: {change_desc}")

        doc.save(str(path))
        logger.info(f"  Saved: {path}")
